from pathlib import Path
from docx import Document
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from xml.sax.saxutils import escape
OUT=Path(__file__).resolve().parent.parent/'output'; OUT.mkdir(exist_ok=True)

def pdf(text,path=OUT/'Jagadeesh-Tailored-CV.pdf'):
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36); styles=getSampleStyleSheet(); story=[]
    for line in text.splitlines():
        if not line: story.append(Spacer(1,6)); continue
        clean=line.replace('**','')
        if clean.startswith('# '): story.append(Paragraph(escape(clean[2:]),styles['Title']))
        elif clean.startswith('## '): story.append(Paragraph(escape(clean[3:]),styles['Heading2']))
        elif clean.startswith('### '): story.append(Paragraph(escape(clean[4:]),styles['Heading3']))
        else: story.append(Paragraph(escape(clean),styles['BodyText']))
    doc.build(story); return path

def docx(text,path=OUT/'Jagadeesh-Tailored-CV.docx'):
    d=Document()
    for line in text.splitlines():
        if line.startswith('# '): d.add_heading(line[2:],0)
        elif line.startswith('## '): d.add_heading(line[3:],1)
        elif line.startswith('### '): d.add_heading(line[4:],2)
        elif line: d.add_paragraph(line.replace('**',''))
    d.save(path); return path
